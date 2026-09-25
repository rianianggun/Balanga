"""Ekspor laporan: Berita Acara Verifikasi (Format Scoring) & Rekap Penilaian per Periode.
Format Excel mengikuti template `data/format_scoring_template.xlsx` (kolom B..G)."""
import io
import re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
SIGN_LEFT = ("Kepala Biro Organisasi Sekretariat Daerah", "Provinsi Kalimantan Tengah,", "BETRI SUSILAWATI, S.Pi", "Pembina Tk. I", "NIP. 19751225 200001 2 001")
SIGN_RIGHT = ("KEPALA ….", "", "(Nama Pejabat)", "(Jenjang)", "NIP. ")


def fmt_tanggal(iso: str) -> str:
    if not iso:
        return "-"
    try:
        d = datetime.fromisoformat(str(iso).replace("Z", "+00:00"))
    except Exception:
        return str(iso)
    return f"{d.day} {BULAN[d.month - 1]} {d.year}"


def fmt_waktu(iso: str) -> str:
    if not iso:
        return "-"
    try:
        d = datetime.fromisoformat(str(iso).replace("Z", "+00:00"))
    except Exception:
        return str(iso)
    return f"{d.strftime('%d/%m/%Y %H:%M')}"


def tempat(area: str) -> str:
    if not area:
        return ""
    if area.lower().startswith("provinsi"):
        return "Palangka Raya"
    return re.sub(r"^(Kota|Kabupaten)\s+", "", area).strip()


def _num(v):
    try:
        f = float(v)
        return int(f) if f.is_integer() else round(f, 2)
    except Exception:
        return v


# ---------------------------------------------------------------------------
# Data assembly for Berita Acara
# ---------------------------------------------------------------------------
def ba_rows(s: dict, indicators: list, link_fn):
    """Return list of row dicts: no, indikator, sebelum, hasil, keterangan, skor."""
    vals = (s.get("scoring") or {}).get("validations") or {}
    uploads = s.get("uploads") or {}
    out = []
    for n, ind in enumerate(indicators, start=1):
        v = vals.get(ind["id"]) or {}
        up = uploads.get(ind["id"]) or {}
        data = v.get("data_validasi") or "-"
        hasil = data if v.get("ok", True) else (f"Tidak valid — {v.get('note')}" if v.get("note") else "Tidak valid")
        ket = "-"
        link = ""
        if up.get("file_id"):
            link = link_fn(up["file_id"]) if link_fn else ""
            ket = up.get("original_filename") or up["file_id"]
        out.append({"no": n, "indikator": ind["name"], "sebelum": data, "hasil": hasil,
                    "keterangan": ket, "link": link, "skor": v.get("score", "")})
    return out


def ba_context(s: dict, umum: list, teknis: list, pengali: float, link_fn):
    sc = s.get("scoring") or {}
    umum_rows = ba_rows(s, umum, link_fn)
    teknis_rows = ba_rows(s, teknis, link_fn)
    umum_total = sc.get("umum_total", sum(float(r["skor"] or 0) for r in umum_rows))
    teknis_total = sc.get("teknis_total", sum(float(r["skor"] or 0) for r in teknis_rows))
    total = round(umum_total + teknis_total, 2)
    final = round(total * pengali, 2)
    urusan = s["urusan"] + (f" — {s['sub_urusan']}" if s.get("sub_urusan") else "")
    area = s["area"]
    narasi = (f"Berdasarkan hasil validasi oleh Biro Organisasi Sekretariat Daerah Provinsi Kalimantan Tengah dan telah dilakukan "
              f"verifikasi data dukung oleh Inspektorat {area} maka skor urusan {urusan} {area} sebelum dikalikan dengan faktor "
              f"kesulitan geografis adalah sebesar {_num(total)} dan skor urusan pemerintahan ini setelah dikalikan dengan faktor "
              f"kesulitan geografis ({_num(pengali)}) adalah {_num(final)}")
    penutup = (f"Demikian Berita Acara Verifikasi Pengisian Data Variabel Pemetaan Urusan Pemerintahan ditandatangani oleh Kepala Biro "
               f"Organisasi Sekretariat Daerah Provinsi Kalimantan Tengah dan Kepala {s['device_name']} {area} untuk dipergunakan sebagaimana mestinya.")
    tgl = fmt_tanggal(sc.get("scored_at") or s.get("updated_at"))
    return {"area": area, "urusan": urusan, "umum": umum_rows, "teknis": teknis_rows, "umum_total": _num(umum_total),
            "teknis_total": _num(teknis_total), "total": _num(total), "pengali": _num(pengali), "final": _num(final),
            "narasi": narasi, "penutup": penutup, "tempat_tgl": f"{tempat(area)}, {tgl}", "device_name": s["device_name"],
            "tahun": s.get("year")}


# ---------------------------------------------------------------------------
# Berita Acara — Excel
# ---------------------------------------------------------------------------
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _table_block(ws, r, letter, title, rows, ctx_total=None):
    """Write one factor table starting at row r. Returns next free row."""
    ws.cell(row=r, column=2, value=letter).font = Font(bold=True)
    ws.cell(row=r, column=3, value=title).font = Font(bold=True)
    r += 1
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
    ws.merge_cells(start_row=r, start_column=7, end_row=r + 1, end_column=7)
    ws.cell(row=r, column=4, value="Data Per Indikator")
    ws.cell(row=r, column=7, value="SKOR")
    heads = ["No", "Indikator", "Data Sebelum Validasi", "Data Hasil Validasi", "Keterangan"]
    for i, h in enumerate(heads):
        ws.cell(row=r + 1, column=2 + i, value=h)
    for rr in (r, r + 1):
        for c in range(2, 8):
            cell = ws.cell(row=rr, column=c)
            cell.font = Font(bold=True); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = BORDER
    r += 2
    for i in range(6):
        cell = ws.cell(row=r, column=2 + i, value=i + 1)
        cell.alignment = Alignment(horizontal="center"); cell.border = BORDER; cell.font = Font(italic=True, size=9)
    r += 1
    for row in rows:
        ws.cell(row=r, column=2, value=f"{row['no']}.")
        ws.cell(row=r, column=3, value=row["indikator"])
        ws.cell(row=r, column=4, value=row["sebelum"])
        ws.cell(row=r, column=5, value=row["hasil"])
        ket = ws.cell(row=r, column=6, value=row["keterangan"])
        if row.get("link"):
            ket.hyperlink = row["link"]; ket.font = Font(color="0563C1", underline="single")
        ws.cell(row=r, column=7, value=_num(row["skor"]) if row["skor"] != "" else "")
        for c in range(2, 8):
            cell = ws.cell(row=r, column=c); cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if c in (2, 4, 5, 7) else "left")
        r += 1
    if ctx_total is not None:
        for label, val in ctx_total:
            ws.cell(row=r, column=6, value=label).font = Font(bold=True)
            ws.cell(row=r, column=7, value=val).font = Font(bold=True)
            for c in (6, 7):
                ws.cell(row=r, column=c).border = BORDER; ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")
            r += 1
    return r + 1


def build_berita_acara_xlsx(ctx: dict) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Berita Acara"
    for col, w in {"A": 4, "B": 5.3, "C": 52.9, "D": 22, "E": 16.3, "F": 36.3, "G": 15}.items():
        ws.column_dimensions[col].width = w
    ws.merge_cells("B1:G1")
    ws["B1"] = "BERITA ACARA VERIFIKASI PENGISIAN DATA VARIABEL PEMETAAN URUSAN PEMERINTAHAN"
    ws["B1"].font = Font(bold=True, size=12); ws["B1"].alignment = Alignment(horizontal="center")
    ws["B3"] = "I. Identitas Daerah dan Urusan Pemerintahan"; ws["B3"].font = Font(bold=True)
    ws["B5"] = "1."; ws["C5"] = "Provinsi"; ws["D5"] = ": KALIMANTAN TENGAH"
    ws["B6"] = "2."; ws["C6"] = "Kabupaten/Kota"; ws["D6"] = f": {ctx['area']}"
    ws["B7"] = "3."; ws["C7"] = "Urusan Pemerintahan"; ws["D7"] = f": {ctx['urusan']}"
    ws["B8"] = "4."; ws["C8"] = "Perangkat Daerah"; ws["D8"] = f": {ctx['device_name']}"
    ws["B9"] = "II. Formulir Validasi Pemetaan Urusan Pemerintahan"; ws["B9"].font = Font(bold=True)
    r = 11
    r = _table_block(ws, r, "A.", "Faktor Umum", ctx["umum"], [("Jumlah", ctx["umum_total"])])
    r = _table_block(ws, r, "B.", "Faktor Teknis", ctx["teknis"],
                     [("Jumlah", ctx["teknis_total"]), ("Jumlah Umum + Teknis", ctx["total"]), ("Pengali", ctx["pengali"]), ("TOTAL", ctx["final"])])
    r += 1
    for text in (ctx["narasi"], ctx["penutup"]):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        c = ws.cell(row=r, column=2, value=text); c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="justify")
        ws.row_dimensions[r].height = 62
        r += 2
    ws.cell(row=r, column=5, value=ctx["tempat_tgl"])
    r += 1
    ws.cell(row=r, column=3, value=SIGN_LEFT[0]); ws.cell(row=r, column=5, value=SIGN_RIGHT[0]); r += 1
    ws.cell(row=r, column=3, value=SIGN_LEFT[1]); r += 4
    for i in (2, 3, 4):
        ws.cell(row=r, column=3, value=SIGN_LEFT[i]).font = Font(bold=(i == 2))
        ws.cell(row=r, column=5, value=SIGN_RIGHT[i]).font = Font(bold=(i == 2))
        r += 1
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Berita Acara — PDF
# ---------------------------------------------------------------------------
def _styles():
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=9, leading=12)
    return {
        "title": ParagraphStyle("t", parent=base, fontName="Helvetica-Bold", fontSize=12, alignment=TA_CENTER, leading=16),
        "h": ParagraphStyle("h", parent=base, fontName="Helvetica-Bold", fontSize=10),
        "base": base,
        "cell": ParagraphStyle("c", parent=base, fontSize=8, leading=10),
        "cellb": ParagraphStyle("cb", parent=base, fontSize=8, leading=10, fontName="Helvetica-Bold", alignment=TA_CENTER),
        "just": ParagraphStyle("j", parent=base, alignment=TA_JUSTIFY),
        "center": ParagraphStyle("ce", parent=base, alignment=TA_CENTER),
        "link": ParagraphStyle("l", parent=base, fontSize=8, leading=10, textColor=colors.HexColor("#0563C1")),
    }


def _pdf_factor_table(letter, title, rows, totals, st):
    data = [[Paragraph(f"<b>{letter} {title}</b>", st["cell"]), "", "", "", "", ""],
            [Paragraph("No", st["cellb"]), Paragraph("Indikator", st["cellb"]), Paragraph("Data Per Indikator", st["cellb"]), "", "", Paragraph("SKOR", st["cellb"])],
            ["", "", Paragraph("Data Sebelum Validasi", st["cellb"]), Paragraph("Data Hasil Validasi", st["cellb"]), Paragraph("Keterangan", st["cellb"]), ""]]
    for r in rows:
        ket = Paragraph(f'<a href="{r["link"]}">{r["keterangan"]}</a>', st["link"]) if r.get("link") else Paragraph(str(r["keterangan"]), st["cell"])
        data.append([Paragraph(f"{r['no']}.", st["cell"]), Paragraph(str(r["indikator"]), st["cell"]), Paragraph(str(r["sebelum"]), st["cell"]),
                     Paragraph(str(r["hasil"]), st["cell"]), ket, Paragraph(str(_num(r["skor"])) if r["skor"] != "" else "", st["cellb"])])
    first_total = len(data)
    for label, val in totals:
        data.append(["", "", "", "", Paragraph(f"<b>{label}</b>", st["cellb"]), Paragraph(f"<b>{_num(val)}</b>", st["cellb"])])
    t = Table(data, colWidths=[1.0 * cm, 7.6 * cm, 3.4 * cm, 3.2 * cm, 6.0 * cm, 1.8 * cm], repeatRows=3)
    style = [("SPAN", (0, 0), (5, 0)), ("SPAN", (2, 1), (4, 1)), ("SPAN", (0, 1), (0, 2)), ("SPAN", (1, 1), (1, 2)), ("SPAN", (5, 1), (5, 2)),
             ("GRID", (0, 1), (-1, first_total - 1), 0.5, colors.black), ("BACKGROUND", (0, 1), (-1, 2), colors.HexColor("#E2E8F0")),
             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("GRID", (4, first_total), (5, -1), 0.5, colors.black)]
    t.setStyle(TableStyle(style))
    return t


def build_berita_acara_pdf(ctx: dict) -> bytes:
    st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            title="Berita Acara Verifikasi", author="Balanga - Biro Organisasi")
    el = [Paragraph("BERITA ACARA VERIFIKASI PENGISIAN DATA VARIABEL PEMETAAN URUSAN PEMERINTAHAN", st["title"]), Spacer(1, 10),
          Paragraph("I. Identitas Daerah dan Urusan Pemerintahan", st["h"]), Spacer(1, 4)]
    ident = Table([["1.", "Provinsi", ": KALIMANTAN TENGAH"], ["2.", "Kabupaten/Kota", f": {ctx['area']}"],
                   ["3.", "Urusan Pemerintahan", f": {ctx['urusan']}"], ["4.", "Perangkat Daerah", f": {ctx['device_name']}"]],
                  colWidths=[0.8 * cm, 4.5 * cm, 17 * cm])
    ident.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("TOPPADDING", (0, 0), (-1, -1), 2)]))
    el += [ident, Spacer(1, 10), Paragraph("II. Formulir Validasi Pemetaan Urusan Pemerintahan", st["h"]), Spacer(1, 6),
           _pdf_factor_table("A.", "Faktor Umum", ctx["umum"], [("Jumlah", ctx["umum_total"])], st), Spacer(1, 10),
           _pdf_factor_table("B.", "Faktor Teknis", ctx["teknis"], [("Jumlah", ctx["teknis_total"]), ("Jumlah Umum + Teknis", ctx["total"]),
                                                                    ("Pengali", ctx["pengali"]), ("TOTAL", ctx["final"])], st),
           Spacer(1, 12), Paragraph(ctx["narasi"], st["just"]), Spacer(1, 8), Paragraph(ctx["penutup"], st["just"]), Spacer(1, 16)]
    sig = Table([["", ctx["tempat_tgl"]], [SIGN_LEFT[0], SIGN_RIGHT[0]], [SIGN_LEFT[1], ""], ["", ""], ["", ""], ["", ""],
                 [SIGN_LEFT[2], SIGN_RIGHT[2]], [SIGN_LEFT[3], SIGN_RIGHT[3]], [SIGN_LEFT[4], SIGN_RIGHT[4]]], colWidths=[11 * cm, 11 * cm])
    sig.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 9), ("FONTNAME", (0, 6), (-1, 6), "Helvetica-Bold"), ("BOTTOMPADDING", (0, 0), (-1, -1), 1), ("TOPPADDING", (0, 0), (-1, -1), 1)]))
    el.append(sig)
    doc.build(el)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Rekap penilaian per periode
# ---------------------------------------------------------------------------
REKAP_HEADERS = ["No", "Kabupaten/Kota", "Perangkat Daerah", "Urusan", "F. Umum", "F. Teknis", "Total", "Nilai Akhir", "Tipe", "Waktu Penilaian", "Penilai"]


def _rekap_matrix(rows):
    out = []
    for i, r in enumerate(rows, start=1):
        out.append([i, r["area"], r["device_name"], r["urusan"] + (f" — {r['sub_urusan']}" if r.get("sub_urusan") else ""),
                    _num(r["umum_total"]), _num(r["teknis_total"]), _num(r["total"]), _num(r["final"]), r["tipe_label"],
                    fmt_waktu(r.get("scored_at")), r.get("penilai_name") or "-"])
    return out


def build_rekap_xlsx(period: dict, rows: list, pengali: float, summary: dict) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Rekap Penilaian"
    head_fill = PatternFill("solid", fgColor="064E3B"); head_font = Font(bold=True, color="FFFFFF")
    ws.append(["REKAP HASIL PENILAIAN TIPOLOGI PERANGKAT DAERAH (PP 18/2016)"]); ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Periode: {period.get('name', '-')} (Tahun {period.get('year', '-')})"])
    ws.append([f"Faktor kesulitan geografis (pengali): {_num(pengali)}"])
    ws.append([f"Dicetak: {fmt_waktu(datetime.utcnow().isoformat())} UTC"])
    ws.append([])
    ws.append(["Ringkasan Tipologi Perangkat Daerah"]); ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for k in ("A", "B", "C", "Lainnya"):
        ws.append([f"Tipe {k}" if k != "Lainnya" else "Lainnya (Bidang/Subbidang)", summary.get(k, 0)])
    ws.append([])
    ws.append(REKAP_HEADERS)
    hr = ws.max_row
    for c in range(1, len(REKAP_HEADERS) + 1):
        cell = ws.cell(row=hr, column=c); cell.fill = head_fill; cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = BORDER
    for row in _rekap_matrix(rows):
        ws.append(row)
        for c in range(1, len(REKAP_HEADERS) + 1):
            cell = ws.cell(row=ws.max_row, column=c); cell.border = BORDER; cell.alignment = Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate([5, 26, 34, 40, 10, 10, 10, 11, 18, 17, 22], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.getvalue()


def build_rekap_pdf(period: dict, rows: list, pengali: float, summary: dict) -> bytes:
    st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=1.2 * cm, bottomMargin=1.2 * cm,
                            title="Rekap Penilaian", author="Balanga - Biro Organisasi")
    el = [Paragraph("REKAP HASIL PENILAIAN TIPOLOGI PERANGKAT DAERAH (PP 18/2016)", st["title"]),
          Paragraph(f"Periode: {period.get('name', '-')} (Tahun {period.get('year', '-')}) · Pengali: {_num(pengali)}", st["center"]), Spacer(1, 8),
          Paragraph(f"Ringkasan: Tipe A = <b>{summary.get('A', 0)}</b> · Tipe B = <b>{summary.get('B', 0)}</b> · Tipe C = <b>{summary.get('C', 0)}</b> · Lainnya = <b>{summary.get('Lainnya', 0)}</b> perangkat daerah", st["base"]),
          Spacer(1, 8)]
    data = [[Paragraph(f"<b>{h}</b>", st["cellb"]) for h in REKAP_HEADERS]]
    for row in _rekap_matrix(rows):
        data.append([Paragraph(str(x), st["cell"]) for x in row])
    if len(data) == 1:
        data.append([Paragraph("Belum ada penilaian selesai pada periode ini.", st["cell"])] + [""] * (len(REKAP_HEADERS) - 1))
    t = Table(data, colWidths=[0.9 * cm, 3.6 * cm, 4.6 * cm, 5.0 * cm, 1.5 * cm, 1.5 * cm, 1.5 * cm, 1.7 * cm, 2.4 * cm, 2.4 * cm, 2.6 * cm], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D1FAE5")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    el.append(t)
    doc.build(el)
    return buf.getvalue()
