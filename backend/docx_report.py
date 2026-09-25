"""Laporan Hasil (Berita Acara) dalam format Word, mengisi template data/format_scoring_template_new.docx."""
import copy
import io
import re
from datetime import datetime, timezone, timedelta
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TEMPLATE = Path(__file__).parent / "data" / "format_scoring_template_new.docx"
HARI = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
SATUAN = ["", "Satu", "Dua", "Tiga", "Empat", "Lima", "Enam", "Tujuh", "Delapan", "Sembilan", "Sepuluh", "Sebelas"]


def terbilang(n: int) -> str:
    n = int(n)
    if n < 12: return SATUAN[n]
    if n < 20: return f"{SATUAN[n - 10]} Belas"
    if n < 100: return f"{SATUAN[n // 10]} Puluh {SATUAN[n % 10]}".strip()
    if n < 200: return f"Seratus {terbilang(n - 100)}".strip()
    if n < 1000: return f"{SATUAN[n // 100]} Ratus {terbilang(n % 100)}".strip()
    if n < 2000: return f"Seribu {terbilang(n - 1000)}".strip()
    return f"{terbilang(n // 1000)} Ribu {terbilang(n % 1000)}".strip()


def _num(v):
    try:
        f = float(v)
        return str(int(f)) if f.is_integer() else f"{f:.2f}".replace(".", ",")
    except Exception:
        return "" if v is None else str(v)


def _set_text(paragraph, text):
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        paragraph.add_run(text)


def _replace_placeholders(paragraph, values: dict):
    if "{" not in paragraph.text:
        return
    def sub(m):
        key = re.sub(r"\s+", " ", m.group(1)).strip().lower()
        return str(values.get(key, m.group(0)))
    new = re.sub(r"\{\s*([^}]+?)\s*\}", sub, paragraph.text)
    if new != paragraph.text:
        _set_text(paragraph, new)


def _add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); rpr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "18"); rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); run.append(t)
    link.append(run); paragraph._p.append(link)


def _fill_cell(cell, text):
    _set_text(cell.paragraphs[0], text)
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def _fill_rows(table, rows, first_data_idx, template_count):
    """Replace `template_count` template rows starting at first_data_idx with len(rows) filled rows."""
    tmpl = table.rows[first_data_idx]
    tmpl_tr = tmpl._tr
    new_trs = []
    for r in rows:
        tr = copy.deepcopy(tmpl_tr)
        new_trs.append(tr)
    # remove template rows
    for i in range(template_count):
        tr = table.rows[first_data_idx]._tr
        tr.getparent().remove(tr)
    anchor = table.rows[first_data_idx - 1]._tr
    for tr in new_trs:
        anchor.addnext(tr); anchor = tr
    from docx.table import _Row
    for i, r in enumerate(rows):
        row = _Row(new_trs[i], table)
        cells = row.cells
        _fill_cell(cells[0], f"{r['no']}.")
        _fill_cell(cells[1], str(r["indikator"]))
        _fill_cell(cells[2], "")
        _fill_cell(cells[3], str(r.get("data_validasi") or ""))
        _fill_cell(cells[4], "")
        if r.get("keterangan") and r.get("keterangan") != "-":
            if r.get("link"):
                _add_hyperlink(cells[4].paragraphs[0], r["link"], str(r["keterangan"]))
            else:
                _fill_cell(cells[4], str(r["keterangan"]))
        _fill_cell(cells[5], _num(r["skor"]) if r["skor"] != "" else "")


def build_berita_acara_docx(ctx: dict, apply_multiplier: bool) -> bytes:
    doc = Document(str(TEMPLATE))
    now = datetime.now(timezone(timedelta(hours=7)))
    total = ctx["total"]; final = ctx["final"] if apply_multiplier else total
    values = {
        "day of today date in indonesian": HARI[now.weekday()], "number of today date": str(now.day),
        "month of today date in indonesian": BULAN[now.month - 1], "year of today date in indonesian words": terbilang(now.year),
        "urusan": ctx["urusan"], "area": ctx["area"], "total": _num(final), "nama perangkat daerah": ctx["device_name"],
    }
    umum_rows = [{**r, "data_validasi": r.get("data_validasi", "")} for r in ctx["umum"]]
    teknis_rows = [{**r, "data_validasi": r.get("data_validasi", "")} for r in ctx["teknis"]]
    t_umum, t_teknis, t_sign = doc.tables[0], doc.tables[1], doc.tables[2]
    _fill_rows(t_umum, umum_rows, 2, 3)
    n_footer = 3
    _fill_rows(t_teknis, teknis_rows, 2, len(t_teknis.rows) - 2 - n_footer)
    footer = t_teknis.rows[-3:]
    _fill_cell(footer[0].cells[-1], _num(total))
    _fill_cell(footer[1].cells[-1], "1,1")
    _fill_cell(footer[2].cells[-1], _num(final))
    if not apply_multiplier:
        tr = footer[1]._tr; tr.getparent().remove(tr)
    for p in doc.paragraphs:
        _replace_placeholders(p, values)
    for row in t_sign.rows:
        for c in row.cells:
            for p in c.paragraphs:
                _replace_placeholders(p, values)
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()
